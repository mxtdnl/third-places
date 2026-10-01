#!/usr/bin/env python3
"""Validate a claim-wall analysis JSON against the schema, the workbook and the theory file.

Usage (from claim-wall/):
    python3 scripts/validate_analysis.py analysis/S1.json
    python3 scripts/validate_analysis.py analysis/S1.json --workbook claim_wall_S1.xlsx

Requires: openpyxl, jsonschema (pip install openpyxl jsonschema).

Exit status 0 when there are no errors (warnings do not fail), 1 otherwise.

Checks, in order:
  1. Schema: the JSON validates against analysis/schema.json (JSON Schema 2020-12).
  2. Theory file: the recorded path resolves, its SHA-256 matches the recorded value
     and the value in ANALYSIS_PROTOCOL.md Section 1, and every characteristic name
     used anywhere in the JSON appears verbatim as a "### N. Name" heading in it.
  3. Workbook coverage: every Entries row for the snapshot is coded exactly once;
     no entry ID in the JSON is absent from the workbook; claim, discipline,
     card text and transcription confidence match the workbook; snapshot counts
     and date match.
  4. Every code is present (schema) and every rationale has at least one quote,
     and every quote is a verbatim substring of that entry's Card_Text.
  5. Every Uncertain entry carries a non-empty transcription caveat.
  6. Internal consistency with ANALYSIS_PROTOCOL.md: lens-fit score 2 implies
     actual_lens.primary equals the filed lens; vocabulary flags for trigger terms;
     claim-synthesis gate, imported candidates, nearest misses, gaps, cluster and
     convergence references; snapshot-wide unengaged characteristics.
"""

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from jsonschema import Draft202012Validator, FormatChecker
except ImportError as exc:  # pragma: no cover
    sys.exit(f"Missing dependency ({exc.name}). Run: pip install openpyxl jsonschema")

ROOT = Path(__file__).resolve().parent.parent  # claim-wall/
LENSES = ("Psychology", "Economics", "Politics")
CLAIM_COLUMNS = ("Entry_ID", "Snapshot_ID", "Snapshot_Date", "Claim_No", "Claim_Label",
                 "Discipline", "Card_Text", "Transcription_Confidence", "Transcription_Notes")
SUBSCORES = ("claim_fit", "mechanism", "explains_change", "lens_fit", "testability", "scope")

# ANALYSIS_PROTOCOL.md Section 2.2. "neutral" on its own needs judgement (only when
# applied to a place, setting or surrounding), so it raises a warning, not an error.
TRIGGERS = {
    "Neutral Ground": [r"\bneutral ground\b"],
    "The Leveler": [r"\blevell?ers?\b", r"\blevell?ing\b", r"\bequali[sz]er\b"],
    "Conversation is the Main Activity": [r"\bconversation"],
    "Accessibility and Accommodation": [r"\baccessib", r"\baccommodat"],
    "The Regulars": [r"\bregulars\b"],
    "A Low Profile": [r"\blow[- ]profile\b"],
    "The Mood is Playful": [r"\bplayful"],
    "A Home Away from Home": [r"\bhome away from home\b"],
}
NEUTRAL_ALONE = r"\bneutral\b(?! ground)"


class Report:
    def __init__(self):
        self.errors, self.warnings, self.passed = [], [], []

    def error(self, msg):
        self.errors.append(msg)

    def warn(self, msg):
        self.warnings.append(msg)

    def check(self, name, before):
        if len(self.errors) == before:
            self.passed.append(name)


def load_workbook(path, snapshot_id):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Entries"]
    rows = list(ws.iter_rows(values_only=True))
    header = tuple(rows[0][:len(CLAIM_COLUMNS)])
    if header != CLAIM_COLUMNS:
        raise SystemExit(f"Unexpected Entries header in {path}: {header}")
    entries = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        rec = dict(zip(CLAIM_COLUMNS, r))
        if rec["Snapshot_ID"] == snapshot_id:
            entries[rec["Entry_ID"]] = rec
    claims = {}
    for r in list(wb["Claims"].iter_rows(values_only=True))[1:]:
        if r[0] is not None:
            claims[int(r[0])] = {"label": r[1], "text": r[2]}
    return entries, claims


def theory_headings(text):
    """Characteristic names: the '### N. Name' headings of Section A only (protocol Section 2)."""
    m = re.search(r"^## A\..*?(?=^## )", text, re.M | re.S)
    if not m:
        raise SystemExit("theory file has no '## A.' section; the characteristic list cannot be located")
    names = set(h.group(1).strip() for h in re.finditer(r"^### \d+\. (.+)$", m.group(0), re.M))
    if len(names) != 8:
        raise SystemExit(f"expected 8 characteristic headings in theory file Section A, found {len(names)}")
    return names


def iter_rationales(entry):
    """Yield (location, rationale) for every rationale object in one entry."""
    eid = entry["entry_id"]
    old = entry["oldenburg"]
    for i, ch in enumerate(old["characteristics"]):
        yield f"{eid} oldenburg.characteristics[{i}] ({ch['name']})", ch["rationale"]
    if "none_rationale" in old:
        yield f"{eid} oldenburg.none_rationale", old["none_rationale"]
    for i, vf in enumerate(old["vocabulary_flags"]):
        yield f"{eid} oldenburg.vocabulary_flags[{i}]", vf["rationale"]
    st = entry["strength"]
    for k in SUBSCORES:
        yield f"{eid} strength.{k}", st[k]["rationale"]
    nae = st["claim_fit"]["not_an_explanation"]
    if nae:
        yield f"{eid} strength.claim_fit.not_an_explanation", nae["rationale"]


def characteristic_names_used(doc):
    names = set()
    for e in doc["entries"]:
        names.update(c["name"] for c in e["oldenburg"]["characteristics"])
        names.update(v["maps_to"] for v in e["oldenburg"]["vocabulary_flags"])
    for cs in doc["claim_synthesis"]:
        names.update(cs["gaps"]["characteristics_unengaged"])
    names.update(doc["snapshot_gaps"]["characteristics_unengaged_anywhere"])
    return names


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("analysis", help="path to analysis/<Snapshot_ID>.json")
    ap.add_argument("--workbook", help="entries workbook (default: snapshot.source_workbook, relative to claim-wall/)")
    ap.add_argument("--schema", default=str(ROOT / "analysis" / "schema.json"))
    ap.add_argument("--protocol", default=str(ROOT / "ANALYSIS_PROTOCOL.md"))
    args = ap.parse_args()

    rep = Report()
    doc = json.loads(Path(args.analysis).read_text(encoding="utf-8"))
    schema = json.loads(Path(args.schema).read_text(encoding="utf-8"))

    # 1. Schema ------------------------------------------------------------
    before = len(rep.errors)
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    for err in sorted(validator.iter_errors(doc), key=lambda e: list(e.absolute_path)):
        loc = "/" + "/".join(str(p) for p in err.absolute_path)
        rep.error(f"schema {loc}: {err.message[:300]}")
    rep.check("JSON validates against schema.json", before)
    if rep.errors:
        return finish(rep)  # later checks assume schema-valid structure

    snap = doc["snapshot"]
    sid = snap["snapshot_id"]
    try:
        datetime.fromisoformat(snap["generated_at"].replace("Z", "+00:00"))
    except ValueError:
        rep.error(f"snapshot.generated_at is not an ISO 8601 date-time: {snap['generated_at']}")

    # 2. Theory file -------------------------------------------------------
    before = len(rep.errors)
    theory_path = (ROOT / doc["theory_file"]["path"]).resolve()
    if not theory_path.is_file():
        rep.error(f"theory file not found at {theory_path}")
        return finish(rep)
    theory_bytes = theory_path.read_bytes()
    digest = hashlib.sha256(theory_bytes).hexdigest()
    if digest != doc["theory_file"]["sha256"]:
        rep.error(f"theory file SHA-256 {digest} != recorded {doc['theory_file']['sha256']}")
    protocol_text = Path(args.protocol).read_text(encoding="utf-8")
    if digest not in protocol_text:
        rep.error("theory file SHA-256 does not match the value recorded in ANALYSIS_PROTOCOL.md "
                  "(the characteristic list may be stale; stop and report)")
    m = re.search(r"\*\*Version:\*\*\s*([0-9][^\s(]*)", protocol_text)
    if m and m.group(1) != doc["protocol_version"]:
        rep.error(f"protocol_version {doc['protocol_version']!r} != ANALYSIS_PROTOCOL.md version {m.group(1)!r}")
    rep.check("theory file hash matches JSON and protocol; protocol version matches", before)

    before = len(rep.errors)
    headings = theory_headings(theory_bytes.decode("utf-8"))
    for name in sorted(characteristic_names_used(doc)):
        if name not in headings:
            rep.error(f"characteristic name {name!r} is not a verbatim '### N. Name' heading in the theory file")
    rep.check("every characteristic name appears verbatim in the theory file", before)

    # 3. Workbook coverage -------------------------------------------------
    before = len(rep.errors)
    wb_path = Path(args.workbook) if args.workbook else ROOT / snap["source_workbook"]
    wb_entries, wb_claims = load_workbook(wb_path, sid)
    if not wb_entries:
        rep.error(f"no Entries rows with Snapshot_ID {sid} in {wb_path}")
        return finish(rep)
    json_ids = [e["entry_id"] for e in doc["entries"]]
    dupes = sorted({i for i in json_ids if json_ids.count(i) > 1})
    if dupes:
        rep.error(f"entry IDs coded more than once: {dupes}")
    missing = sorted(set(wb_entries) - set(json_ids))
    extra = sorted(set(json_ids) - set(wb_entries))
    if missing:
        rep.error(f"workbook entries with no code in JSON: {missing}")
    if extra:
        rep.error(f"entry IDs in JSON absent from workbook snapshot {sid}: {extra}")
    if json_ids != sorted(json_ids):
        rep.warn("entries are not in Entry_ID order (protocol Section 4)")
    rep.check("every workbook entry is coded once; no JSON entry is absent from the workbook", before)

    before = len(rep.errors)
    entries = {e["entry_id"]: e for e in doc["entries"] if e["entry_id"] in wb_entries}
    for eid, e in entries.items():
        w = wb_entries[eid]
        if e["card_text"] != w["Card_Text"]:
            rep.error(f"{eid} card_text differs from workbook Card_Text")
        if e["claim_no"] != w["Claim_No"]:
            rep.error(f"{eid} claim_no {e['claim_no']} != workbook {w['Claim_No']}")
        if e["discipline"] != w["Discipline"]:
            rep.error(f"{eid} discipline {e['discipline']} != workbook {w['Discipline']}")
        if e["transcription_confidence"] != w["Transcription_Confidence"]:
            rep.error(f"{eid} transcription_confidence differs from workbook")
    if snap["entry_count"] != len(wb_entries):
        rep.error(f"snapshot.entry_count {snap['entry_count']} != workbook {len(wb_entries)}")
    n_unc = sum(1 for w in wb_entries.values() if w["Transcription_Confidence"] == "Uncertain")
    if snap["uncertain_count"] != n_unc:
        rep.error(f"snapshot.uncertain_count {snap['uncertain_count']} != workbook {n_unc}")
    wb_dates = {w["Snapshot_Date"] for w in wb_entries.values()}
    if len(wb_dates) > 1:
        rep.error(f"workbook has more than one Snapshot_Date for {sid}: {wb_dates}")
    wb_date = next(iter(wb_dates))
    wb_date = wb_date.date().isoformat() if hasattr(wb_date, "date") else wb_date
    if snap["snapshot_date"] != wb_date:
        rep.error(f"snapshot.snapshot_date {snap['snapshot_date']!r} != workbook {wb_date!r}")
    rep.check("card text, claim, discipline, confidence, counts and date match the workbook", before)

    # 4. Quotes ------------------------------------------------------------
    before = len(rep.errors)
    n_quotes = 0
    for eid, e in entries.items():
        text = wb_entries[eid]["Card_Text"]
        for loc, rat in iter_rationales(e):
            if not rat["quotes"]:
                rep.error(f"{loc}: rationale has no quote")
            for q in rat["quotes"]:
                n_quotes += 1
                if q not in text:
                    rep.error(f"{loc}: quote not in Card_Text: {q!r}")
                if "[" in q or "]" in q:
                    rep.error(f"{loc}: quote contains transcriber brackets: {q!r}")
            if not rat["text"].strip():
                rep.error(f"{loc}: empty rationale text")
        for i, vf in enumerate(e["oldenburg"]["vocabulary_flags"]):
            if vf["term_as_written"] not in text:
                rep.error(f"{eid} vocabulary_flags[{i}]: term_as_written {vf['term_as_written']!r} not in Card_Text")
    rep.check(f"every rationale quotes Card_Text verbatim ({n_quotes} quotes)", before)

    # 5. Uncertain entries -------------------------------------------------
    before = len(rep.errors)
    for eid, e in entries.items():
        cav = e["transcription_caveat"]
        if e["transcription_confidence"] == "Uncertain":
            if not cav or not cav.strip():
                rep.error(f"{eid} is Uncertain but has no transcription_caveat")
            if not e["caveat_affects"]:
                rep.warn(f"{eid} is Uncertain but caveat_affects is empty (protocol 6.1)")
        elif not cav and e["caveat_affects"]:
            rep.error(f"{eid} lists caveat_affects but has no caveat")
    rep.check("every Uncertain entry carries a caveat", before)

    # 6. Protocol consistency ----------------------------------------------
    before = len(rep.errors)
    for eid, e in entries.items():
        st, old = e["strength"], e["oldenburg"]
        lf = st["lens_fit"]
        if lf["score"] == 2 and lf["actual_lens"]["primary"] != e["discipline"]:
            rep.error(f"{eid} lens_fit 2 but actual_lens.primary {lf['actual_lens']['primary']} != filed {e['discipline']}")
        if lf["score"] == 0 and lf["actual_lens"]["primary"] == e["discipline"]:
            rep.error(f"{eid} lens_fit 0 but actual_lens.primary equals the filed lens")
        if lf["actual_lens"]["primary"] in lf["actual_lens"]["secondary"]:
            rep.error(f"{eid} actual_lens.primary repeated in secondary")
        mf = st["claim_fit"]["misfiled_to"]
        if mf is not None and mf == e["claim_no"]:
            rep.error(f"{eid} misfiled_to equals the filed claim")
        names = [c["name"] for c in old["characteristics"]]
        if len(names) != len(set(names)):
            rep.error(f"{eid} maps the same characteristic more than once")
        if old["engagement"] == "engaged" and "none_rationale" in old:
            rep.error(f"{eid} engaged but carries none_rationale")
        sc = st["scope"]
        if sc["score"] == 0 and sc["stated_scope"]:
            rep.warn(f"{eid} scope 0 but stated_scope is filled")
        if sc["score"] > 0 and not sc["stated_scope"]:
            rep.error(f"{eid} scope {sc['score']} but stated_scope is empty")
        # Vocabulary flags for trigger terms (protocol 2.2, 7.3)
        text_l = e["card_text"].lower()
        flagged = {v["maps_to"] for v in old["vocabulary_flags"]}
        for ch, pats in TRIGGERS.items():
            if any(re.search(p, text_l) for p in pats) and ch not in flagged:
                rep.error(f"{eid} contains a trigger term for {ch!r} but has no vocabulary flag")
        if re.search(NEUTRAL_ALONE, text_l) and "Neutral Ground" not in flagged:
            rep.warn(f"{eid} uses 'neutral' without a Neutral Ground flag; check whether it applies to a place")
        for v in old["vocabulary_flags"]:
            if v["maps_to"] not in names:
                rep.warn(f"{eid} flags vocabulary for {v['maps_to']!r} but does not map that characteristic")
    rep.check("entry codes are internally consistent (lens, misfiling, scope, vocabulary flags)", before)

    before = len(rep.errors)
    all_ids = set(entries)

    def gate(e):
        s = e["strength"]
        return s["claim_fit"]["score"] == 2 and s["mechanism"]["score"] >= 1 and s["explains_change"]["score"] >= 1

    def gate_fails(e):
        s = e["strength"]
        out = []
        if s["claim_fit"]["score"] != 2:
            out.append("claim_fit")
        if s["mechanism"]["score"] < 1:
            out.append("mechanism")
        if s["explains_change"]["score"] < 1:
            out.append("explains_change")
        return out

    def subscores(e):
        return {k: e["strength"][k]["score"] for k in SUBSCORES}

    claim_nos = [cs["claim_no"] for cs in doc["claim_synthesis"]]
    if sorted(claim_nos) != [1, 2, 3, 4]:
        rep.error(f"claim_synthesis must cover claims 1-4 once each; got {claim_nos}")
    clusters_by_id, engaged_anywhere = {}, set()
    for e in entries.values():
        engaged_anywhere.update(c["name"] for c in e["oldenburg"]["characteristics"])

    for cs in doc["claim_synthesis"]:
        n = cs["claim_no"]
        tag = f"claim_synthesis[{n}]"
        if n in wb_claims:
            if cs["claim_label"] != wb_claims[n]["label"]:
                rep.error(f"{tag} claim_label != workbook Claims sheet")
            if cs["claim_text"] != wb_claims[n]["text"]:
                rep.error(f"{tag} claim_text != workbook Claims sheet")
        filed = [e for e in entries.values() if e["claim_no"] == n]

        expected = {e["entry_id"] for e in filed if gate(e)}
        listed = [c["entry_id"] for c in cs["candidates"]]
        if set(listed) != expected:
            rep.error(f"{tag} candidates {sorted(listed)} != entries passing the 9.1 gate {sorted(expected)}")
        ranks = [tuple(entries[i]["strength"][k]["score"] for k in ("mechanism", "explains_change", "testability"))
                 for i in listed if i in entries]
        if ranks != sorted(ranks, reverse=True):
            rep.error(f"{tag} candidates not ordered by mechanism, then explains_change, then testability")

        imported = {e["entry_id"] for e in entries.values() if e["strength"]["claim_fit"]["misfiled_to"] == n}
        listed_imp = [c["entry_id"] for c in cs["imported_candidates"]]
        if set(listed_imp) != imported:
            rep.error(f"{tag} imported_candidates {sorted(listed_imp)} != entries with misfiled_to={n} {sorted(imported)}")

        for c in cs["candidates"] + cs["imported_candidates"]:
            if c["entry_id"] in entries and c["subscores"] != subscores(entries[c["entry_id"]]):
                rep.error(f"{tag} subscores for {c['entry_id']} do not match the entry's codes")

        if not cs["candidates"] and not cs["nearest_misses"]:
            rep.error(f"{tag} has no candidates and no nearest_misses (protocol 9.1)")
        for nm in cs["nearest_misses"]:
            e = entries.get(nm["entry_id"])
            if e is None:
                rep.error(f"{tag} nearest_miss {nm['entry_id']} is not a coded entry")
                continue
            if e["claim_no"] != n:
                rep.error(f"{tag} nearest_miss {nm['entry_id']} is not filed under this claim")
            if sorted(nm["fails"]) != sorted(gate_fails(e)):
                rep.error(f"{tag} nearest_miss {nm['entry_id']} fails {nm['fails']} but codes give {gate_fails(e)}")

        for cl in cs["clusters"]:
            ids = set(cl["entry_ids"])
            if not ids <= all_ids:
                rep.error(f"{tag} cluster {cl['cluster_id']} cites unknown entries {sorted(ids - all_ids)}")
            if not any(entries[i]["claim_no"] == n for i in ids & all_ids):
                rep.error(f"{tag} cluster {cl['cluster_id']} has no member filed under this claim")
            q_ids = {q["entry_id"] for q in cl["quotes"]}
            if len(q_ids) < 2:
                rep.error(f"{tag} cluster {cl['cluster_id']} quotes fewer than two cards")
            for q in cl["quotes"]:
                if q["entry_id"] not in ids:
                    rep.error(f"{tag} cluster {cl['cluster_id']} quotes {q['entry_id']}, which is not a member")
                elif q["entry_id"] in entries and q["quote"] not in entries[q["entry_id"]]["card_text"]:
                    rep.error(f"{tag} cluster {cl['cluster_id']} quote not in {q['entry_id']} Card_Text: {q['quote']!r}")
            prev = clusters_by_id.setdefault(cl["cluster_id"], cl)
            if prev is not cl and (set(prev["entry_ids"]) != ids or prev["kind"] != cl["kind"]):
                rep.error(f"cluster {cl['cluster_id']} is defined differently under different claims")

        for i, cv in enumerate(cs["convergence"]):
            ids = set(cv["entry_ids"])
            if not ids <= all_ids:
                rep.error(f"{tag} convergence[{i}] cites unknown entries {sorted(ids - all_ids)}")
                continue
            if not any(entries[j]["claim_no"] == n for j in ids):
                rep.error(f"{tag} convergence[{i}] has no member filed under this claim")
            avail = set()
            for j in ids:
                al = entries[j]["strength"]["lens_fit"]["actual_lens"]
                avail.update([entries[j]["discipline"], al["primary"], *al["secondary"]])
            if not set(cv["lenses"]) <= avail:
                rep.error(f"{tag} convergence[{i}] lists lenses {cv['lenses']} not filed or coded for its entries")
            filed_lenses = {entries[j]["discipline"] for j in ids}
            if len(filed_lenses) < 2:
                rep.warn(f"{tag} convergence[{i}] entries are all filed under one lens ({filed_lenses})")

        for i, tn in enumerate(cs["tensions"]):
            ids = set(tn["entry_ids"])
            if not ids <= all_ids:
                rep.error(f"{tag} tensions[{i}] cites unknown entries {sorted(ids - all_ids)}")
            elif not any(entries[j]["claim_no"] == n for j in ids):
                rep.error(f"{tag} tensions[{i}] has no member filed under this claim")

        engaged = set()
        for e in filed:
            engaged.update(c["name"] for c in e["oldenburg"]["characteristics"])
        exp_unengaged = headings - engaged
        if set(cs["gaps"]["characteristics_unengaged"]) != exp_unengaged:
            rep.error(f"{tag} characteristics_unengaged {sorted(cs['gaps']['characteristics_unengaged'])} "
                      f"!= computed {sorted(exp_unengaged)}")
        with_mech = {e["strength"]["lens_fit"]["actual_lens"]["primary"] for e in filed
                     if e["strength"]["mechanism"]["score"] >= 1}
        exp_lwm = {l for l in LENSES if l not in with_mech}
        if set(cs["gaps"]["lenses_without_mechanism"]) != exp_lwm:
            rep.error(f"{tag} lenses_without_mechanism {sorted(cs['gaps']['lenses_without_mechanism'])} "
                      f"!= computed {sorted(exp_lwm)}")

    exp_any = headings - engaged_anywhere
    if set(doc["snapshot_gaps"]["characteristics_unengaged_anywhere"]) != exp_any:
        rep.error(f"snapshot_gaps.characteristics_unengaged_anywhere != computed {sorted(exp_any)}")
    rep.check("claim synthesis is consistent with the entry codes (gate, imports, misses, clusters, gaps)", before)

    if doc["overrides"]:
        rep.warn(f"{len(doc['overrides'])} override(s) present; analyst-written files should have none")

    return finish(rep, doc)


def finish(rep, doc=None):
    for p in rep.passed:
        print(f"PASS  {p}")
    for w in rep.warnings:
        print(f"WARN  {w}")
    for e in rep.errors:
        print(f"FAIL  {e}")
    if doc is not None and not rep.errors:
        flags = summarise(doc)
        print()
        print(flags)
    print()
    print(f"{len(rep.errors)} error(s), {len(rep.warnings)} warning(s)")
    return 1 if rep.errors else 0


def summarise(doc):
    """List entries flagged misfiled_to, not_an_explanation, or actual lens != filed lens."""
    lines = ["Flagged entries (spec Session 2 'Done when'):"]
    for e in doc["entries"]:
        cf = e["strength"]["claim_fit"]
        al = e["strength"]["lens_fit"]["actual_lens"]
        bits = []
        if cf["misfiled_to"] is not None:
            bits.append(f"misfiled_to={cf['misfiled_to']}")
        if cf["not_an_explanation"]:
            bits.append(f"not_an_explanation={cf['not_an_explanation']['type']}")
        if al["primary"] != e["discipline"]:
            bits.append(f"actual_lens={al['primary']} (filed {e['discipline']})")
        if bits:
            lines.append(f"  {e['entry_id']}  claim {e['claim_no']}  " + "; ".join(bits))
    return "\n".join(lines)


if __name__ == "__main__":
    sys.exit(main())
