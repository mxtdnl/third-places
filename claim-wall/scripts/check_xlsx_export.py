#!/usr/bin/env python3
"""Check a viewer .xlsx export against the matching viewer JSON export.

Usage, from claim-wall/:
    python3 scripts/check_xlsx_export.py <export.xlsx> <export.json>

The JSON export holds the analyst's codes unchanged plus the `overrides`
array. This script applies the overrides itself (RFC 6901 pointers,
independently of the viewer's code) to get the reconciled document, opens the
.xlsx with openpyxl, and checks that:
  - the four sheets exist (Entries_Reconciled, Entries_Analyst, Overrides, Snapshot)
  - every Entries_Reconciled cell equals the reconciled JSON value
  - every Entries_Analyst cell equals the analyst's JSON value
  - Overrides has one row per override, with original and new values equal
    to the JSON (compared as parsed JSON, not as text)
  - Snapshot holds the snapshot ID and override count
Exit 0 when every check passes. Needs openpyxl (pip install openpyxl).
"""
import copy
import json
import sys

import openpyxl

CRITERIA = [
    ("claim_fit", "Claim_fit"), ("mechanism", "Mechanism"), ("explains_change", "Explains_change"),
    ("lens_fit", "Lens_fit"), ("testability", "Testability"), ("scope", "Scope"),
]


def ptr_tokens(p):
    return [t.replace("~1", "/").replace("~0", "~") for t in p[1:].split("/")]


def ptr_get(doc, p):
    cur = doc
    for t in ptr_tokens(p):
        cur = cur[int(t)] if isinstance(cur, list) else cur[t]
    return cur


def ptr_set(doc, p, v):
    toks = ptr_tokens(p)
    cur = doc
    for t in toks[:-1]:
        cur = cur[int(t)] if isinstance(cur, list) else cur[t]
    last = toks[-1]
    if isinstance(cur, list):
        cur[int(last)] = v
    else:
        cur[last] = v


def rationale_text(r):
    return " | ".join(f'"{q}"' for q in r["quotes"]) + " — " + r["text"]


def blank(v):
    # The writer leaves None and "" as empty cells; openpyxl reads them as None.
    return None if v in (None, "") else v


def expected_row(e, doc):
    labels = {c["claim_no"]: c["claim_label"] for c in doc["claim_synthesis"]}
    o, s = e["oldenburg"], e["strength"]
    row = {
        "Entry_ID": e["entry_id"],
        "Claim_No": e["claim_no"],
        "Claim_Label": labels[e["claim_no"]],
        "Discipline": e["discipline"],
        "Card_Text": e["card_text"],
        "Transcription_Confidence": e["transcription_confidence"],
        "Transcription_Caveat": e["transcription_caveat"],
        "Caveat_Affects": ", ".join(e["caveat_affects"]),
        "Oldenburg_Engagement": o["engagement"],
        "Oldenburg_Characteristics": "; ".join(f'{c["name"]} ({c["relation"]})' for c in o["characteristics"]),
        "Oldenburg_Rationales": "\n".join(
            [f'{c["name"]}: {rationale_text(c["rationale"])}' for c in o["characteristics"]]
            + ([f'None: {rationale_text(o["none_rationale"])}'] if o.get("none_rationale") else [])),
        "Vocabulary_Flags": "\n".join(
            f'"{f["term_as_written"]}" -> {f["maps_to"]} ({f["usage_matches_definition"]}): {rationale_text(f["rationale"])}'
            for f in o["vocabulary_flags"]),
    }
    for key, title in CRITERIA:
        row[f"{title}_Score"] = s[key]["score"]
        row[f"{title}_Rationale"] = rationale_text(s[key]["rationale"])
    nae = s["claim_fit"]["not_an_explanation"]
    row.update({
        "Misfiled_To": s["claim_fit"]["misfiled_to"],
        "Not_An_Explanation": f'{nae["type"]}: {rationale_text(nae["rationale"])}' if nae else None,
        "Actual_Lens_Primary": s["lens_fit"]["actual_lens"]["primary"],
        "Actual_Lens_Secondary": ", ".join(s["lens_fit"]["actual_lens"]["secondary"]),
        "Testable_Implication": s["testability"]["testable_implication"],
        "Implication_Note": s["testability"]["implication_note"],
        "Stated_Scope": s["scope"]["stated_scope"],
        "Bodo_Question": s["scope"]["bodo_question"],
        "Analyst_Note": e.get("analyst_note"),
    })
    return {k: blank(v) for k, v in row.items()}


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    return head, [dict(zip(head, r)) for r in rows[1:]]


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1])
    doc = json.load(open(sys.argv[2], encoding="utf-8"))
    rec = copy.deepcopy(doc)
    for ov in doc["overrides"]:
        ptr_set(rec, ov["pointer"], copy.deepcopy(ov["override_value"]))

    errors, checked = [], 0

    def eq(where, got, want):
        nonlocal checked
        checked += 1
        if got != want:
            errors.append(f"{where}: xlsx {got!r} != json {want!r}")

    want_sheets = ["Entries_Reconciled", "Entries_Analyst", "Overrides", "Snapshot"]
    eq("sheet names", wb.sheetnames, want_sheets)
    if errors:
        print("\n".join(errors))
        return 1

    overridden = {}
    for ov in doc["overrides"]:
        overridden.setdefault(ov["target"]["entry_id"], []).append(ov)

    for name, source in (("Entries_Reconciled", rec), ("Entries_Analyst", doc)):
        head, rows = sheet_rows(wb[name])
        eq(f"{name} row count", len(rows), len(source["entries"]))
        for e, r in zip(source["entries"], rows):
            for col, want in expected_row(e, source).items():
                eq(f"{name} {e['entry_id']} {col}", r.get(col), want)
            if name == "Entries_Reconciled":
                eq(f"{name} {e['entry_id']} Overridden_Fields non-empty", bool(r.get("Overridden_Fields")), e["entry_id"] in overridden)
        extra = set(head) - set(expected_row(source["entries"][0], source)) - {"Overridden_Fields"}
        eq(f"{name} unexpected columns", sorted(extra), [])

    head, rows = sheet_rows(wb["Overrides"])
    eq("Overrides row count", len(rows), len(doc["overrides"]))
    for ov, r in zip(doc["overrides"], rows):
        tag = f"Overrides {ov['override_id']}"
        eq(f"{tag} Override_ID", r["Override_ID"], ov["override_id"])
        eq(f"{tag} Entry_ID", r["Entry_ID"], ov["target"]["entry_id"])
        eq(f"{tag} Pointer", r["Pointer"], ov["pointer"])
        eq(f"{tag} Original_JSON", json.loads(r["Original_JSON"]), ov["original_value"])
        eq(f"{tag} Override_JSON", json.loads(r["Override_JSON"]), ov["override_value"])
        eq(f"{tag} original is the analyst's value", ov["original_value"], ptr_get(doc, ov["pointer"]))
        eq(f"{tag} Reason", r["Reason"], ov["reason"])
        eq(f"{tag} By", r["By"], ov["by"])
        eq(f"{tag} At", r["At"], ov["at"])

    _, rows = sheet_rows(wb["Snapshot"])
    meta = {r["Field"]: r["Value"] for r in rows}
    eq("Snapshot Snapshot_ID", meta.get("Snapshot_ID"), doc["snapshot"]["snapshot_id"])
    eq("Snapshot Override_Count", meta.get("Override_Count"), len(doc["overrides"]))
    eq("Snapshot Entry_Count", meta.get("Entry_Count"), doc["snapshot"]["entry_count"])

    for e in errors[:30]:
        print("FAIL ", e)
    print(f"{checked} cell checks, {len(errors)} failures ({len(doc['overrides'])} overrides, {len(doc['entries'])} entries)")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
